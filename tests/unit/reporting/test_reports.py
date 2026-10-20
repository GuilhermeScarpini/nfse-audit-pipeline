import json
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from nfse_audit.reporting import BatchReport, DocumentError, to_excel, to_json
from nfse_audit.synthetic.builders import make_invoice, make_iss


def _report(engine) -> BatchReport:
    return BatchReport(
        batch_id="batch-1",
        generated_at=datetime(2026, 10, 20, tzinfo=UTC),
        invoices=[
            engine.audit(make_invoice()),
            engine.audit(make_invoice(access_key="2" * 50, iss=make_iss(amount="40.00"))),
        ],
        errors=[DocumentError(source="broken.xml", message="not an NFS-e")],
    )


def test_summary_counts(engine):
    assert _report(engine).summary == {
        "documents": 3,
        "invoices_audited": 2,
        "documents_with_errors": 1,
        "invoices_ok": 1,
        "invoices_divergent": 1,
        "invoices_error": 0,
    }


def test_json_keeps_decimals_as_strings(engine):
    data = json.loads(to_json(_report(engine)))
    iss_amount = next(r for r in data["invoices"][1]["results"] if r["rule_id"] == "ISS_AMOUNT")
    assert (iss_amount["expected"], iss_amount["found"], iss_amount["status"]) == (
        "50.00",
        "40.00",
        "divergent",
    )
    assert data["summary"]["invoices_divergent"] == 1


def test_excel_has_one_detail_row_per_rule_check(engine):
    report = _report(engine)
    workbook = load_workbook(BytesIO(to_excel(report)))
    assert workbook.sheetnames == ["Summary", "Invoices", "Details", "Errors"]
    checks = sum(len(audit.results) for audit in report.invoices)
    assert workbook["Details"].max_row == 1 + checks
    assert workbook["Invoices"].max_row == 3
    assert workbook["Errors"]["A2"].value == "broken.xml"


def test_excel_keeps_amounts_numeric(engine):
    workbook = load_workbook(BytesIO(to_excel(_report(engine))))
    amount = workbook["Invoices"]["F2"]
    assert isinstance(amount.value, int | float)
    assert amount.number_format == "#,##0.00####"
