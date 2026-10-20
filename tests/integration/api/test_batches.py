from io import BytesIO
from typing import ClassVar

import fakeredis
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from nfse_audit.api import dependencies
from nfse_audit.api.main import create_app
from nfse_audit.config import Settings, get_settings
from nfse_audit.domain import Invoice
from nfse_audit.extraction import ExtractionError, ExtractorRegistry
from nfse_audit.storage import BatchFileStore, BatchStatusStore
from nfse_audit.synthetic.builders import make_invoice, make_iss
from nfse_audit.worker.tasks import run_batch


class FakeExtractor:
    name = "fake"
    invoices: ClassVar[dict[bytes, Invoice]] = {
        b"ok": make_invoice(),
        b"bad": make_invoice(access_key="2" * 50, iss=make_iss(amount="1.00")),
    }

    def supports(self, filename: str, content: bytes) -> bool:
        return True

    def extract(self, content: bytes) -> Invoice:
        try:
            return self.invoices[content]
        except KeyError:
            raise ExtractionError("unreadable") from None


class InlineQueue:
    """Runs the job immediately, in-process, instead of sending it to a worker."""

    def __init__(self, status_store, file_store, engine, run: bool = True) -> None:
        self.enqueued: list[str] = []
        self._args = (status_store, file_store, ExtractorRegistry([FakeExtractor()]), engine)
        self._run = run

    def enqueue_batch(self, batch_id: str) -> None:
        self.enqueued.append(batch_id)
        if self._run:
            run_batch(batch_id, *self._args)


@pytest.fixture
def setup(tmp_path, engine):
    settings = Settings(data_dir=tmp_path, max_files_per_batch=3, max_file_bytes=100)
    status_store = BatchStatusStore(fakeredis.FakeRedis())
    file_store = BatchFileStore(tmp_path)
    queue = InlineQueue(status_store, file_store, engine)
    app = create_app()
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[dependencies.get_status_store] = lambda: status_store
    app.dependency_overrides[dependencies.get_file_store] = lambda: file_store
    app.dependency_overrides[dependencies.get_queue] = lambda: queue
    return TestClient(app), queue, status_store


def _files(*contents: bytes):
    return [
        ("files", (f"nota{i}.xml", content, "application/xml"))
        for i, content in enumerate(contents)
    ]


def test_full_batch_flow(setup):
    client, queue, _ = setup
    response = client.post("/batches", files=_files(b"ok", b"bad", b"???"))
    assert response.status_code == 202
    body = response.json()
    batch_id = body["batch_id"]
    assert queue.enqueued == [batch_id]
    assert body["status_url"].endswith(f"/batches/{batch_id}")

    batch = client.get(f"/batches/{batch_id}").json()
    assert batch["state"] == "done"
    assert batch["summary"]["invoices_divergent"] == 1
    assert batch["summary"]["documents_with_errors"] == 1

    report = client.get(f"/batches/{batch_id}/report").json()
    assert len(report["invoices"]) == 2

    xlsx = client.get(f"/batches/{batch_id}/report", params={"format": "xlsx"})
    assert xlsx.headers["content-type"].startswith("application/vnd.openxml")
    assert load_workbook(BytesIO(xlsx.content)).sheetnames[0] == "Summary"


def test_idempotency_key_returns_the_same_batch(setup):
    client, queue, _ = setup
    headers = {"Idempotency-Key": "upload-42"}
    first = client.post("/batches", files=_files(b"ok"), headers=headers).json()
    second = client.post("/batches", files=_files(b"ok"), headers=headers).json()
    assert first["batch_id"] == second["batch_id"]
    assert queue.enqueued == [first["batch_id"]]  # processed only once


def test_report_of_unfinished_batch_is_409(setup, engine, tmp_path):
    client, queue, _ = setup
    queue._run = False
    batch_id = client.post("/batches", files=_files(b"ok")).json()["batch_id"]
    assert client.get(f"/batches/{batch_id}").json()["state"] == "queued"
    assert client.get(f"/batches/{batch_id}/report").status_code == 409


def test_unknown_batch_is_404(setup):
    client, _, _ = setup
    assert client.get("/batches/" + "0" * 32).status_code == 404


@pytest.mark.parametrize(
    "files",
    [
        _files(b"ok", b"ok", b"ok", b"ok"),  # more than max_files_per_batch
        _files(b"x" * 101),  # larger than max_file_bytes
    ],
)
def test_limits(setup, files):
    client, queue, _ = setup
    assert client.post("/batches", files=files).status_code == 413
    assert queue.enqueued == []


def test_health(setup):
    client, _, _ = setup
    assert client.get("/health").json() == {"status": "ok"}
