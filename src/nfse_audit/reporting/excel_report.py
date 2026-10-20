"""Excel report: summary, one row per invoice, one row per rule check, and errors."""

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from nfse_audit.domain.cnpj import format_cnpj
from nfse_audit.reporting.models import BatchReport
from nfse_audit.rules.results import RuleStatus

STATUS_FILL = {
    RuleStatus.OK: PatternFill("solid", fgColor="C6EFCE"),
    RuleStatus.DIVERGENT: PatternFill("solid", fgColor="FFC7CE"),
    RuleStatus.ERROR: PatternFill("solid", fgColor="FFEB9C"),
    RuleStatus.NOT_APPLICABLE: PatternFill("solid", fgColor="EDEDED"),
}
VALUE_FORMAT = "#,##0.00####"
HEADER_FONT = Font(bold=True)


def to_excel(report: BatchReport) -> bytes:
    workbook = Workbook()
    _summary(workbook.active, report)
    _invoices(workbook.create_sheet("Invoices"), report)
    _details(workbook.create_sheet("Details"), report)
    _errors(workbook.create_sheet("Errors"), report)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _summary(sheet: Worksheet, report: BatchReport) -> None:
    sheet.title = "Summary"
    rows: list[list[Any]] = [
        ["Batch", report.batch_id],
        ["Generated at", report.generated_at.isoformat(timespec="seconds")],
        *[[key.replace("_", " ").capitalize(), value] for key, value in report.summary.items()],
    ]
    for row in rows:
        sheet.append(row)
        sheet.cell(sheet.max_row, 1).font = HEADER_FONT
    _autosize(sheet)


def _invoices(sheet: Worksheet, report: BatchReport) -> None:
    _header(
        sheet,
        [
            "Access key",
            "Number",
            "Competence",
            "Provider",
            "Taker",
            "Amount",
            "Status",
            "Divergent",
        ],
    )
    for audit in report.invoices:
        divergent = sum(r.status is RuleStatus.DIVERGENT for r in audit.results)
        sheet.append(
            [
                audit.access_key,
                audit.number,
                audit.competence_date,
                format_cnpj(audit.provider_cnpj),
                format_cnpj(audit.taker_cnpj),
                audit.service_amount,
                audit.status.value,
                divergent,
            ]
        )
        _style_row(sheet, status=audit.status, status_col=7, value_cols=(6,), date_cols=(3,))
    _finish(sheet)


def _details(sheet: Worksheet, report: BatchReport) -> None:
    _header(
        sheet,
        [
            "Access key",
            "Competence",
            "Rule",
            "Description",
            "Expected",
            "Found",
            "Status",
            "Detail",
            "Legal basis",
            "Table version",
            "Table verified",
        ],
    )
    for audit in report.invoices:
        for result in audit.results:
            sheet.append(
                [
                    audit.access_key,
                    audit.competence_date,
                    result.rule_id,
                    result.description,
                    result.expected,
                    result.found,
                    result.status.value,
                    result.detail,
                    result.legal_basis,
                    result.table_version,
                    _yes_no(result.table_verified),
                ]
            )
            _style_row(sheet, status=result.status, status_col=7, value_cols=(5, 6), date_cols=(2,))
    _finish(sheet)


def _errors(sheet: Worksheet, report: BatchReport) -> None:
    _header(sheet, ["Source", "Message"])
    for error in report.errors:
        sheet.append([error.source, error.message])
    _finish(sheet)


def _header(sheet: Worksheet, titles: list[str]) -> None:
    sheet.append(titles)
    for cell in sheet[1]:
        cell.font = HEADER_FONT


def _style_row(
    sheet: Worksheet,
    *,
    status: RuleStatus,
    status_col: int,
    value_cols: tuple[int, ...],
    date_cols: tuple[int, ...],
) -> None:
    row = sheet.max_row
    sheet.cell(row, status_col).fill = STATUS_FILL[status]
    for col in value_cols:
        sheet.cell(row, col).number_format = VALUE_FORMAT
    for col in date_cols:
        sheet.cell(row, col).number_format = "yyyy-mm-dd"


def _finish(sheet: Worksheet) -> None:
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions
    _autosize(sheet)


def _autosize(sheet: Worksheet) -> None:
    for column in sheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width + 2, 60)


def _yes_no(value: bool | None) -> str | None:
    return None if value is None else ("yes" if value else "no")
